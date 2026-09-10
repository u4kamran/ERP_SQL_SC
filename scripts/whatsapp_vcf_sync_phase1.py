"""
PHASE 1 — WhatsApp VCF → CUST_SMS synchronization ANALYSIS (READ-ONLY).

ABSOLUTE SAFETY:
  - Only SELECT against the business database.
  - Never INSERT / UPDATE / DELETE / MERGE / ALTER / DROP.
  - Does not modify the VCF file.
  - Writes report files only under data/SC_WhatsappNo_Import/reports/.

Matching rule (existing project convention):
  mobile_key = last 10 digits of digits-only phone (national number).
  Valid Pakistan mobile: 10 digits starting with '3'.
  Storage format for NEW inserts (Phase 2 later): 03XXXXXXXXX via _mobile_for_storage.
"""

from __future__ import annotations

import csv
import json
import re
import sys
from collections import defaultdict
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import text

from app.database.business_session import business_engine

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_VCF = (
    ROOT
    / "data"
    / "SC_WhatsappNo_Import"
    / "Police Helpline and 256 other contacts.vcf"
)
REPORT_DIR = ROOT / "data" / "SC_WhatsappNo_Import" / "reports"


@dataclass
class VcfContact:
    row_num: int
    fn: str = ""
    n_family: str = ""
    n_given: str = ""
    org: str = ""
    title: str = ""
    note: str = ""
    email: str = ""
    phones_raw: list[str] = field(default_factory=list)
    preferred_phone_raw: str = ""
    mobile_key: str = ""
    storage_phone: str = ""
    waid: str = ""
    raw_block: str = ""


def normalize_name(value: str) -> str:
    text_val = (value or "").strip().lower()
    text_val = re.sub(r"[^\w\s]", " ", text_val, flags=re.UNICODE)
    text_val = re.sub(r"\s+", " ", text_val).strip()
    return text_val


def names_effectively_same(a: str, b: str) -> bool:
    return normalize_name(a) == normalize_name(b) and bool(normalize_name(a))


def mobile_key(value: str) -> str:
    digits = re.sub(r"\D", "", value or "")
    if digits.startswith("00"):
        digits = digits[2:]
    if len(digits) >= 10:
        return digits[-10:]
    return ""


def is_valid_pk_mobile_key(key: str) -> bool:
    return len(key) == 10 and key.startswith("3") and key.isdigit()


def mobile_for_storage(value: str) -> str:
    """Project habit: store as local 03XXXXXXXXX when possible."""
    digits = re.sub(r"\D", "", value or "")
    if not digits:
        return ""
    if digits.startswith("00"):
        digits = digits[2:]
    if digits.startswith("92") and len(digits) >= 12:
        return "0" + digits[2:12]
    if len(digits) >= 11 and digits.startswith("0"):
        return digits[:11]
    if len(digits) >= 10:
        return "0" + digits[-10:]
    return digits


def unfold_vcf(content: str) -> str:
    # RFC 6350 line unfolding: CRLF + space/tab continuation
    content = content.replace("\r\n", "\n").replace("\r", "\n")
    return re.sub(r"\n[ \t]", "", content)


def parse_n_field(value: str) -> tuple[str, str]:
    # N:Family;Given;Additional;Prefix;Suffix
    parts = value.split(";")
    family = (parts[0] if len(parts) > 0 else "").strip()
    given = (parts[1] if len(parts) > 1 else "").strip()
    return family, given


def decode_quoted_printable(value: str) -> str:
    if "=E" not in value.upper() and "=C" not in value.upper() and "= " not in value:
        # Still try if soft line breaks / hex escapes present
        if "=" not in value:
            return value
    try:
        import quopri

        return quopri.decodestring(value.encode("utf-8", errors="ignore")).decode(
            "utf-8", errors="replace"
        )
    except Exception:
        return value


def unescape_vcf_text(value: str) -> str:
    value = value.replace("\\n", "\n").replace("\\,", ",").replace("\\;", ";")
    return value.strip()


def parse_vcf(path: Path) -> list[VcfContact]:
    raw = path.read_text(encoding="utf-8", errors="replace")
    content = unfold_vcf(raw)
    blocks = re.findall(
        r"BEGIN:VCARD(.*?)END:VCARD", content, flags=re.IGNORECASE | re.DOTALL
    )
    contacts: list[VcfContact] = []
    for idx, block in enumerate(blocks, start=1):
        c = VcfContact(row_num=idx, raw_block=block.strip())
        phones: list[tuple[str, str, bool]] = []  # (raw, type, is_wa)
        for line in block.split("\n"):
            line = line.strip()
            if not line or ":" not in line:
                continue
            left, right = line.split(":", 1)
            left_u = left.upper()
            prop = left_u.split(";")[0]
            params = left_u
            right = unescape_vcf_text(right)
            if "ENCODING=QUOTED-PRINTABLE" in params:
                right = decode_quoted_printable(right)

            if prop == "FN":
                c.fn = right
            elif prop == "N":
                c.n_family, c.n_given = parse_n_field(right)
            elif prop == "ORG":
                c.org = right.split(";")[0].strip()
            elif prop == "TITLE":
                c.title = right
            elif prop == "NOTE":
                c.note = right
            elif prop == "EMAIL":
                if not c.email:
                    c.email = right
            elif prop == "TEL":
                waid_m = re.search(r"WAID=([^;:]+)", params, re.I)
                waid = waid_m.group(1) if waid_m else ""
                if waid and not c.waid:
                    c.waid = re.sub(r"\D", "", waid)
                is_mobile = "CELL" in params or "MOBILE" in params or "TYPE=CELL" in params
                phones.append((right, "mobile" if is_mobile else "tel", bool(waid)))

        # Prefer WhatsApp-tagged / mobile / first usable
        preferred = ""
        if phones:
            ranked = sorted(
                phones,
                key=lambda p: (
                    0 if p[2] else 1,
                    0 if p[1] == "mobile" else 1,
                    0 if is_valid_pk_mobile_key(mobile_key(p[0])) else 1,
                ),
            )
            preferred = ranked[0][0]
            # If waid looks like a full PK number, prefer it when TEL is weak
            if c.waid and is_valid_pk_mobile_key(mobile_key(c.waid)):
                if not is_valid_pk_mobile_key(mobile_key(preferred)):
                    preferred = c.waid if c.waid.startswith("+") else f"+{c.waid}"

        c.phones_raw = [p[0] for p in phones]
        c.preferred_phone_raw = preferred
        key = mobile_key(preferred)
        if (not is_valid_pk_mobile_key(key)) and c.waid:
            key = mobile_key(c.waid)
            if is_valid_pk_mobile_key(key) and not preferred:
                c.preferred_phone_raw = f"+{c.waid}" if not c.waid.startswith("+") else c.waid
        c.mobile_key = key if is_valid_pk_mobile_key(key) else ""
        c.storage_phone = mobile_for_storage(c.preferred_phone_raw or c.waid) if c.mobile_key else ""
        if not c.fn:
            built = f"{c.n_given} {c.n_family}".strip()
            c.fn = built or c.org or c.title or ""
        contacts.append(c)
    return contacts


def display_name(c: VcfContact) -> str:
    name = (c.fn or "").strip()
    if name:
        return name[:150]
    built = f"{c.n_given} {c.n_family}".strip()
    if built:
        return built[:150]
    if c.org:
        return c.org[:150]
    if c.title:
        return c.title[:150]
    return ""


def name_score(name: str) -> tuple[int, int, int]:
    """Higher is better for picking preferred VCF duplicate."""
    n = (name or "").strip()
    return (len(n), 0 if n.isupper() else 1, n.count(" "))


def load_cust_sms_index() -> tuple[dict[str, list[dict[str, Any]]], list[dict[str, Any]]]:
    """READ-ONLY: SELECT all CUST_SMS phones into mobile_key index."""
    sql = text(
        """
        SELECT
            CUST_ID AS cust_id,
            RTRIM(ISNULL(CUST_NAME, '')) AS cust_name,
            RTRIM(ISNULL(MOBILE_NO, '')) AS mobile_no,
            RTRIM(ISNULL(MOBILE_NO_TMP, '')) AS mobile_no_tmp,
            RTRIM(ISNULL(CUST_ADDRESS, '')) AS cust_address,
            STATUS AS status,
            ADDED_DATETIME AS added_datetime
        FROM dbo.CUST_SMS
        ORDER BY CUST_ID
        """
    )
    with business_engine.connect() as conn:
        rows = [dict(r) for r in conn.execute(sql).mappings().all()]

    index: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        keys: set[str] = set()
        for field_name in ("mobile_no", "mobile_no_tmp"):
            key = mobile_key(str(row.get(field_name) or ""))
            if is_valid_pk_mobile_key(key):
                keys.add(key)
        for key in keys:
            index[key].append(row)
    return index, rows


def classify(
    contacts: list[VcfContact],
    db_index: dict[str, list[dict[str, Any]]],
) -> list[dict[str, Any]]:
    # Detect VCF duplicates by mobile_key
    by_key: dict[str, list[VcfContact]] = defaultdict(list)
    for c in contacts:
        if c.mobile_key:
            by_key[c.mobile_key].append(c)

    preferred_for_key: dict[str, int] = {}
    for key, group in by_key.items():
        if len(group) > 1:
            best = max(group, key=lambda x: name_score(display_name(x)))
            preferred_for_key[key] = best.row_num

    results: list[dict[str, Any]] = []
    for c in contacts:
        vcf_name = display_name(c)
        phone_raw = c.preferred_phone_raw
        key = c.mobile_key

        row: dict[str, Any] = {
            "row_num": c.row_num,
            "vcf_name": vcf_name,
            "vcf_phone": phone_raw,
            "normalized_phone": key,
            "storage_phone": c.storage_phone,
            "vcf_org": c.org,
            "vcf_title": c.title,
            "vcf_email": c.email,
            "vcf_note": (c.note or "")[:200],
            "all_phones": " | ".join(c.phones_raw),
            "waid": c.waid,
            "db_cust_id": None,
            "db_name": None,
            "db_mobile": None,
            "db_match_count": 0,
            "db_all_ids": "",
            "match_status": "",
            "category": "",
            "recommended_action": "",
            "reason": "",
            "is_preferred_vcf_duplicate": False,
        }

        if not key:
            row.update(
                {
                    "match_status": "INVALID",
                    "category": "D",
                    "recommended_action": "SKIP",
                    "reason": "No usable Pakistan mobile (need 10 digits starting with 3).",
                }
            )
            results.append(row)
            continue

        vcf_dupes = by_key.get(key, [])
        if len(vcf_dupes) > 1:
            is_pref = preferred_for_key.get(key) == c.row_num
            row["is_preferred_vcf_duplicate"] = is_pref
            other_names = [
                display_name(x) for x in vcf_dupes if x.row_num != c.row_num
            ]
            if not is_pref:
                row.update(
                    {
                        "match_status": "DUPLICATE_VCF",
                        "category": "E",
                        "recommended_action": "SKIP (VCF DUPLICATE)",
                        "reason": (
                            f"Same normalized phone appears {len(vcf_dupes)} times in VCF. "
                            f"Preferred row #{preferred_for_key.get(key)}. "
                            f"Other names: {', '.join(other_names[:5])}"
                        ),
                    }
                )
                results.append(row)
                continue
            # Preferred VCF duplicate continues into A/B/C/F classification
            row["reason"] = (
                f"Preferred among {len(vcf_dupes)} VCF duplicates "
                f"(others: {', '.join(other_names[:5])}). "
            )

        matches = db_index.get(key, [])
        row["db_match_count"] = len(matches)
        if matches:
            row["db_all_ids"] = ", ".join(str(m["cust_id"]) for m in matches)

        if len(matches) > 1:
            # Prefer lowest CUST_ID for display, but do NOT auto-update
            primary = sorted(matches, key=lambda m: int(m["cust_id"]))[0]
            row["db_cust_id"] = int(primary["cust_id"])
            row["db_name"] = primary["cust_name"]
            row["db_mobile"] = primary["mobile_no"]
            row.update(
                {
                    "match_status": "DATABASE_DUPLICATE",
                    "category": "F",
                    "recommended_action": "MANUAL REVIEW",
                    "reason": row["reason"]
                    + (
                        f"Normalized phone matches {len(matches)} CUST_SMS rows "
                        f"({row['db_all_ids']}). Do not auto-update."
                    ),
                }
            )
            results.append(row)
            continue

        if len(matches) == 1:
            m = matches[0]
            row["db_cust_id"] = int(m["cust_id"])
            row["db_name"] = m["cust_name"]
            row["db_mobile"] = m["mobile_no"]
            same = names_effectively_same(vcf_name, m["cust_name"] or "")
            if same or (not vcf_name.strip()):
                row.update(
                    {
                        "match_status": "SAME",
                        "category": "A",
                        "recommended_action": "NO CHANGE",
                        "reason": row["reason"]
                        + (
                            "Phone exists; name effectively the same."
                            if same
                            else "Phone exists; VCF name empty — keep DB name."
                        ),
                    }
                )
            else:
                row.update(
                    {
                        "match_status": "NAME_CHANGE",
                        "category": "B",
                        "recommended_action": "UPDATE NAME",
                        "reason": row["reason"]
                        + f"Phone exists; VCF name differs from DB ('{m['cust_name']}').",
                    }
                )
            results.append(row)
            continue

        # No DB match
        if not vcf_name.strip():
            row.update(
                {
                    "match_status": "NEW_NO_NAME",
                    "category": "D",
                    "recommended_action": "SKIP",
                    "reason": row["reason"]
                    + "Phone not in DB but VCF name is empty — cannot insert.",
                }
            )
        else:
            row.update(
                {
                    "match_status": "NEW",
                    "category": "C",
                    "recommended_action": "INSERT",
                    "reason": row["reason"] + "Phone not found in CUST_SMS.",
                }
            )
        results.append(row)

    return results


def summarize(results: list[dict[str, Any]]) -> dict[str, int]:
    def count(cat: str) -> int:
        return sum(1 for r in results if r["category"] == cat)

    return {
        "total_vcf_contacts": len(results),
        "valid_contacts": sum(1 for r in results if r["normalized_phone"]),
        "invalid_contacts": count("D"),
        "duplicate_vcf_contacts": count("E"),
        "existing_same_name": count("A"),
        "existing_name_change": count("B"),
        "new_customers": count("C"),
        "database_duplicates": count("F"),
        "requires_manual_review": count("F")
        + sum(
            1
            for r in results
            if r["category"] == "E" and r.get("is_preferred_vcf_duplicate")
        ),
        "no_change": count("A"),
        "names_to_update": count("B"),
        "inserts": count("C"),
    }


def write_csv(path: Path, results: list[dict[str, Any]]) -> None:
    cols = [
        "row_num",
        "vcf_name",
        "vcf_phone",
        "normalized_phone",
        "storage_phone",
        "db_cust_id",
        "db_name",
        "db_mobile",
        "db_match_count",
        "db_all_ids",
        "match_status",
        "category",
        "recommended_action",
        "reason",
        "vcf_org",
        "vcf_title",
        "vcf_email",
        "all_phones",
        "waid",
        "is_preferred_vcf_duplicate",
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        w.writeheader()
        for r in results:
            w.writerow(r)


def write_excel(
    path: Path,
    results: list[dict[str, Any]],
    summary: dict[str, int],
    meta: dict[str, Any],
) -> None:
    wb = Workbook()

    # Summary sheet
    ws = wb.active
    ws.title = "Summary"
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    ws.append(["WhatsApp Customer Contact Sync — Phase 1 Analysis (READ-ONLY)"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Generated: {meta['generated_at']}"])
    ws.append([f"VCF File: {meta['vcf_path']}"])
    ws.append([f"Database: {meta['database']}"])
    ws.append([f"Target Table: {meta['table']}"])
    ws.append([f"Matching Key: {meta['matching_key']}"])
    ws.append([f"CUST_SMS Row Count (at analysis): {meta['db_row_count']}"])
    ws.append([])
    ws.append(["Metric", "Count"])
    ws["A9"].fill = header_fill
    ws["B9"].fill = header_fill
    ws["A9"].font = header_font
    ws["B9"].font = header_font
    for k, v in summary.items():
        ws.append([k, v])
    ws.append([])
    ws.append(["IMPORTANT", "NO database changes were made during Phase 1."])
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 80

    # Detail sheet
    detail = wb.create_sheet("Detailed Report")
    cols = [
        ("row_num", "Row #"),
        ("vcf_name", "VCF Name"),
        ("vcf_phone", "VCF Phone"),
        ("normalized_phone", "Normalized Phone"),
        ("storage_phone", "Storage Phone (if insert)"),
        ("db_cust_id", "Database Customer ID"),
        ("db_name", "Database Name"),
        ("db_mobile", "Database Mobile"),
        ("db_match_count", "DB Match Count"),
        ("db_all_ids", "All DB IDs"),
        ("match_status", "Match Status"),
        ("category", "Category"),
        ("recommended_action", "Recommended Action"),
        ("reason", "Reason"),
        ("vcf_org", "ORG"),
        ("vcf_title", "Title"),
        ("vcf_email", "Email"),
        ("all_phones", "All Phones"),
        ("waid", "WAID"),
    ]
    detail.append([c[1] for c in cols])
    for cell in detail[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True)
    fills = {
        "A": PatternFill("solid", fgColor="C6EFCE"),
        "B": PatternFill("solid", fgColor="FFEB9C"),
        "C": PatternFill("solid", fgColor="BDD7EE"),
        "D": PatternFill("solid", fgColor="D9D9D9"),
        "E": PatternFill("solid", fgColor="F4B183"),
        "F": PatternFill("solid", fgColor="FFC7CE"),
    }
    for r in results:
        detail.append([r.get(c[0]) for c in cols])
        cat = r.get("category") or ""
        fill = fills.get(cat)
        if fill:
            for cell in detail[detail.max_row]:
                cell.fill = fill
    for col in detail.columns:
        detail.column_dimensions[col[0].column_letter].width = min(36, max(12, len(str(col[0].value or "")) + 4))

    # Change preview sheets
    updates = wb.create_sheet("Preview UPDATES")
    updates.append(
        ["Customer ID", "Phone", "OLD NAME", "NEW NAME", "Normalized Phone", "Reason"]
    )
    for cell in updates[1]:
        cell.fill = header_fill
        cell.font = header_font
    for r in results:
        if r["category"] == "B":
            updates.append(
                [
                    r["db_cust_id"],
                    r["db_mobile"] or r["vcf_phone"],
                    r["db_name"],
                    r["vcf_name"],
                    r["normalized_phone"],
                    r["reason"],
                ]
            )

    inserts = wb.create_sheet("Preview INSERTS")
    inserts.append(
        ["VCF Name", "Phone (VCF)", "Storage Phone", "Normalized Phone", "ORG", "Reason"]
    )
    for cell in inserts[1]:
        cell.fill = header_fill
        cell.font = header_font
    for r in results:
        if r["category"] == "C":
            inserts.append(
                [
                    r["vcf_name"],
                    r["vcf_phone"],
                    r["storage_phone"],
                    r["normalized_phone"],
                    r["vcf_org"],
                    r["reason"],
                ]
            )

    nochange = wb.create_sheet("Preview NO CHANGE")
    nochange.append(["Customer ID", "Phone", "Name", "Action"])
    for cell in nochange[1]:
        cell.fill = header_fill
        cell.font = header_font
    for r in results:
        if r["category"] == "A":
            nochange.append(
                [r["db_cust_id"], r["db_mobile"] or r["vcf_phone"], r["db_name"], "NO CHANGE"]
            )

    review = wb.create_sheet("Manual Review")
    review.append(
        [
            "Row #",
            "Category",
            "VCF Name",
            "Phone",
            "Normalized",
            "DB IDs",
            "Action",
            "Reason",
        ]
    )
    for cell in review[1]:
        cell.fill = header_fill
        cell.font = header_font
    for r in results:
        if r["category"] in {"D", "E", "F"}:
            review.append(
                [
                    r["row_num"],
                    r["category"],
                    r["vcf_name"],
                    r["vcf_phone"],
                    r["normalized_phone"],
                    r["db_all_ids"],
                    r["recommended_action"],
                    r["reason"],
                ]
            )

    schema = wb.create_sheet("Schema Mapping")
    schema.append(["Table", "Column", "Purpose", "Matching / Role"])
    for cell in schema[1]:
        cell.fill = header_fill
        cell.font = header_font
    for row in [
        ["dbo.CUST_SMS", "CUST_ID", "Primary key (identity)", "Customer ID"],
        ["dbo.CUST_SMS", "CUST_NAME", "Customer display name", "Name field (update target)"],
        ["dbo.CUST_SMS", "MOBILE_NO", "Primary mobile", "Phone field (match source)"],
        ["dbo.CUST_SMS", "MOBILE_NO_TMP", "Alternate/legacy mobile", "Also matched"],
        ["dbo.CUST_SMS", "CUST_ADDRESS", "Address", "NOT updated in this sync"],
        ["dbo.CUST_SMS", "STATUS", "Status flag", "NOT updated; INSERT uses 1"],
        ["dbo.CUST_SMS", "ADDED_DATETIME", "Created timestamp", "SET on INSERT only"],
        [
            "Matching Key",
            "last 10 digits",
            "Canonical comparison key",
            "Project convention (import/delivery/bot/cart)",
        ],
    ]:
        schema.append(row)

    wb.save(path)


def write_change_preview_txt(path: Path, results: list[dict[str, Any]]) -> None:
    lines: list[str] = []
    lines.append("=" * 72)
    lines.append("CHANGE PREVIEW — Phase 1 (NO DATABASE WRITES)")
    lines.append("=" * 72)

    lines.append("\n### UPDATE (name only)\n")
    updates = [r for r in results if r["category"] == "B"]
    if not updates:
        lines.append("(none)\n")
    for r in updates:
        lines.append(f"Customer ID: {r['db_cust_id']}")
        lines.append(f"Phone: {r['db_mobile'] or r['vcf_phone']}")
        lines.append("OLD NAME:")
        lines.append(str(r["db_name"]))
        lines.append("NEW NAME:")
        lines.append(str(r["vcf_name"]))
        lines.append("-" * 40)

    lines.append("\n### INSERT\n")
    inserts = [r for r in results if r["category"] == "C"]
    if not inserts:
        lines.append("(none)\n")
    for r in inserts:
        lines.append(f"Name: {r['vcf_name']}")
        lines.append(f"Phone (VCF): {r['vcf_phone']}")
        lines.append(f"Phone (storage): {r['storage_phone']}")
        lines.append(f"Normalized: {r['normalized_phone']}")
        lines.append("-" * 40)

    lines.append("\n### NO CHANGE\n")
    for r in [x for x in results if x["category"] == "A"][:50]:
        lines.append(f"Phone: {r['db_mobile'] or r['vcf_phone']}")
        lines.append(f"Name: {r['db_name']}")
        lines.append("Action: NO CHANGE")
        lines.append("-" * 20)
    more = max(0, sum(1 for x in results if x["category"] == "A") - 50)
    if more:
        lines.append(f"... and {more} more NO CHANGE rows (see Excel/CSV).\n")

    lines.append("\n### SKIP / MANUAL REVIEW\n")
    for r in results:
        if r["category"] in {"D", "E", "F"}:
            lines.append(
                f"[{r['category']}] Row {r['row_num']}: {r['vcf_name']} | "
                f"{r['vcf_phone']} | {r['recommended_action']} | {r['reason']}"
            )

    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> int:
    vcf_path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_VCF
    if not vcf_path.exists():
        print(f"ERROR: VCF not found: {vcf_path}")
        return 1

    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    base = f"WhatsApp_Customer_Sync_Report_{stamp}"

    print("=" * 72)
    print("PHASE 1 — READ-ONLY ANALYSIS")
    print("=" * 72)
    print(f"VCF: {vcf_path}")
    print("Parsing VCF...")
    contacts = parse_vcf(vcf_path)
    print(f"Parsed {len(contacts)} vCard entries.")

    print("Loading CUST_SMS (SELECT only)...")
    db_index, all_rows = load_cust_sms_index()
    print(f"CUST_SMS rows: {len(all_rows)}")
    print(f"Distinct valid mobile keys in DB: {len(db_index)}")
    db_dup_keys = sum(1 for k, v in db_index.items() if len(v) > 1)
    print(f"DB keys with >1 row (pre-existing duplicates): {db_dup_keys}")

    print("Classifying...")
    results = classify(contacts, db_index)
    summary = summarize(results)

    meta = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "vcf_path": str(vcf_path),
        "database": "nsds2626 (business)",
        "table": "dbo.CUST_SMS",
        "matching_key": "last 10 digits of digits-only phone (PK mobile starting with 3)",
        "db_row_count": len(all_rows),
        "dry_run": True,
        "database_writes": False,
        "phase": 1,
    }

    xlsx_path = REPORT_DIR / f"{base}.xlsx"
    csv_path = REPORT_DIR / f"{base}.csv"
    json_path = REPORT_DIR / f"{base}.json"
    preview_path = REPORT_DIR / f"{base}_CHANGE_PREVIEW.txt"
    summary_path = REPORT_DIR / f"{base}_SUMMARY.txt"

    write_excel(xlsx_path, results, summary, meta)
    write_csv(csv_path, results)
    write_change_preview_txt(preview_path, results)

    payload = {
        "meta": meta,
        "summary": summary,
        "results": results,
        "approval_required": True,
        "approval_phrases": [
            "YES, SYNC",
            "APPROVE SYNCHRONIZATION",
            "SYNC NEW ONLY",
            "SYNC NAME UPDATES ONLY",
            "SYNC ALL APPROVED",
        ],
    }
    json_path.write_text(json.dumps(payload, indent=2, ensure_ascii=False, default=str), encoding="utf-8")

    summary_lines = [
        "Synchronization Analysis Complete",
        f"Total VCF Contacts: {summary['total_vcf_contacts']}",
        f"Valid Contacts: {summary['valid_contacts']}",
        f"Invalid Contacts: {summary['invalid_contacts']}",
        f"Duplicate VCF Contacts: {summary['duplicate_vcf_contacts']}",
        f"No Change (Same Name): {summary['existing_same_name']}",
        f"Names to Update: {summary['existing_name_change']}",
        f"New Customers: {summary['new_customers']}",
        f"Database Duplicates: {summary['database_duplicates']}",
        f"Requires Manual Review: {summary['requires_manual_review']}",
        "",
        "Database changes have NOT been made.",
        f"Excel: {xlsx_path}",
        f"CSV:   {csv_path}",
        f"JSON:  {json_path}",
        f"Preview: {preview_path}",
    ]
    summary_text = "\n".join(summary_lines)
    summary_path.write_text(summary_text, encoding="utf-8")
    print()
    print(summary_text)
    print()
    print("Do you want me to synchronize the approved records?")
    print("Reply with one of: YES, SYNC | APPROVE SYNCHRONIZATION | SYNC NEW ONLY |")
    print("                   SYNC NAME UPDATES ONLY | SYNC ALL APPROVED")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
