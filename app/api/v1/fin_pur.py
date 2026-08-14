"""Purchase Receipt API — VB6 Fin_PurM / Fin_PurD parity."""

from typing import List, Optional

from fastapi import APIRouter, Depends, File, Query, UploadFile
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.api.deps import CurrentUser, require_permission
from app.config.settings import settings
from app.database.business_session import get_business_db
from app.database.session import get_db
from app.models.user import UserPreference
from app.schemas.fin_pur import (
    FinPurConfigOut,
    FinPurDeleteResponse,
    FinPurDocumentOut,
    FinPurItemOut,
    FinPurLookupRow,
    FinPurSaveRequest,
    FinPurSaveResponse,
    FinPurSupplierOut,
)
from app.services.fin_pur_service import FinPurService

router = APIRouter()


def _legacy_uid(db: Session, user_id: int) -> int:
    row = db.execute(
        select(UserPreference).where(
            UserPreference.UserId == user_id,
            UserPreference.PreferenceKey == "legacy_contpl_uid",
            UserPreference.IsDeleted == False,  # noqa: E712
        )
    ).scalar_one_or_none()
    if row and row.PreferenceValue:
        try:
            return int(row.PreferenceValue)
        except ValueError:
            pass
    return settings.voucher_legacy_uid


def _service(
    current_user: CurrentUser,
    business_db: Session,
    auth_db: Session,
) -> FinPurService:
    return FinPurService(
        business_db,
        username=current_user.username,
        legacy_uid=_legacy_uid(auth_db, current_user.user_id),
        is_admin=current_user.has_permission("auth.admin.full"),
        is_sa=current_user.username.upper() == "SA" or current_user.has_permission("auth.admin.full"),
    )


@router.get("/config", response_model=FinPurConfigOut)
def get_config(
    current_user: CurrentUser = Depends(require_permission("inventory.fin_pur.view")),
    business_db: Session = Depends(get_business_db),
    auth_db: Session = Depends(get_db),
):
    return _service(current_user, business_db, auth_db).get_config()


@router.get("/documents/{prod_id}", response_model=FinPurDocumentOut)
def load_document(
    prod_id: int,
    current_user: CurrentUser = Depends(require_permission("inventory.fin_pur.view")),
    business_db: Session = Depends(get_business_db),
    auth_db: Session = Depends(get_db),
):
    return _service(current_user, business_db, auth_db).load_document(prod_id)


@router.post("/documents", response_model=FinPurSaveResponse)
def save_document(
    payload: FinPurSaveRequest,
    current_user: CurrentUser = Depends(require_permission("inventory.fin_pur.view")),
    business_db: Session = Depends(get_business_db),
    auth_db: Session = Depends(get_db),
):
    from fastapi import HTTPException, status

    is_edit = bool(payload.prod_id and payload.prod_id > 0)
    if is_edit:
        if not (
            current_user.has_permission("inventory.fin_pur.edit")
            or current_user.has_permission("auth.admin.full")
        ):
            raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Insufficient permissions.")
    else:
        if not (
            current_user.has_permission("inventory.fin_pur.create")
            or current_user.has_permission("auth.admin.full")
        ):
            raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Insufficient permissions.")
    return _service(current_user, business_db, auth_db).save_document(payload)


@router.delete("/documents/{prod_id}", response_model=FinPurDeleteResponse)
def delete_document(
    prod_id: int,
    current_user: CurrentUser = Depends(require_permission("inventory.fin_pur.delete")),
    business_db: Session = Depends(get_business_db),
    auth_db: Session = Depends(get_db),
):
    return _service(current_user, business_db, auth_db).delete_document(prod_id)


@router.get("/suppliers/{supplier_id}", response_model=FinPurSupplierOut)
def get_supplier(
    supplier_id: int,
    current_user: CurrentUser = Depends(require_permission("inventory.fin_pur.view")),
    business_db: Session = Depends(get_business_db),
    auth_db: Session = Depends(get_db),
):
    return _service(current_user, business_db, auth_db).get_supplier(supplier_id)


@router.get("/suppliers", response_model=List[FinPurLookupRow])
def search_suppliers(
    q: str = Query(..., min_length=1),
    current_user: CurrentUser = Depends(require_permission("inventory.fin_pur.view")),
    business_db: Session = Depends(get_business_db),
    auth_db: Session = Depends(get_db),
):
    return _service(current_user, business_db, auth_db).search_suppliers(q)


@router.get("/items/by-id/{item_id}", response_model=FinPurItemOut)
def item_by_id(
    item_id: float,
    current_user: CurrentUser = Depends(require_permission("inventory.fin_pur.view")),
    business_db: Session = Depends(get_business_db),
    auth_db: Session = Depends(get_db),
):
    return _service(current_user, business_db, auth_db).get_item(item_id=item_id)


@router.get("/items/by-manual/{manual_id}", response_model=FinPurItemOut)
def item_by_manual(
    manual_id: float,
    current_user: CurrentUser = Depends(require_permission("inventory.fin_pur.view")),
    business_db: Session = Depends(get_business_db),
    auth_db: Session = Depends(get_db),
):
    return _service(current_user, business_db, auth_db).get_item(manual_id=manual_id)


@router.get("/items/by-barcode/{barcode}", response_model=FinPurItemOut)
def item_by_barcode(
    barcode: str,
    current_user: CurrentUser = Depends(require_permission("inventory.fin_pur.view")),
    business_db: Session = Depends(get_business_db),
    auth_db: Session = Depends(get_db),
):
    return _service(current_user, business_db, auth_db).get_item(barcode=barcode)


@router.get("/items/by-ws-barcode/{barcode}", response_model=FinPurItemOut)
def item_by_ws_barcode(
    barcode: str,
    current_user: CurrentUser = Depends(require_permission("inventory.fin_pur.view")),
    business_db: Session = Depends(get_business_db),
    auth_db: Session = Depends(get_db),
):
    return _service(current_user, business_db, auth_db).get_item(barcode_ws=barcode)


@router.get("/items/search", response_model=List[FinPurLookupRow])
def search_items(
    q: str = Query(..., min_length=1),
    current_user: CurrentUser = Depends(require_permission("inventory.fin_pur.view")),
    business_db: Session = Depends(get_business_db),
    auth_db: Session = Depends(get_db),
):
    return _service(current_user, business_db, auth_db).search_items(q)


class AssignCoRequest(BaseModel):
    co_id: int
    item_ids: List[float] = Field(default_factory=list)


@router.post("/assign-co")
def assign_co(
    payload: AssignCoRequest,
    current_user: CurrentUser = Depends(require_permission("inventory.fin_pur.edit")),
    business_db: Session = Depends(get_business_db),
    auth_db: Session = Depends(get_db),
):
    return _service(current_user, business_db, auth_db).assign_company(payload.co_id, payload.item_ids)


class DetailItemUpdate(BaseModel):
    item_id: float
    mrp: float = 0
    stax_rate: float = 0
    stax_amt: float = 0
    qty: float
    amount: float = 0
    disc_amt: float = 0
    off_inv_disc_amt: float = 0
    rate: float = 0
    co_id: int = 0


@router.post("/detail-item-update")
def detail_item_update(
    payload: DetailItemUpdate,
    current_user: CurrentUser = Depends(require_permission("inventory.fin_pur.create")),
    business_db: Session = Depends(get_business_db),
    auth_db: Session = Depends(get_db),
):
    _service(current_user, business_db, auth_db).update_item_from_detail(**payload.model_dump())
    return {"message": "ok"}


@router.post("/excel/clear-qty")
async def excel_clear_qty(
    file: UploadFile = File(...),
    current_user: CurrentUser = Depends(require_permission("inventory.fin_pur.create")),
):
    """Clear Qty column on Sheet 1 Column C for rows with Manual ID in Column A — VB6 cmdClearQty."""
    from io import BytesIO

    from fastapi import HTTPException, status
    from fastapi.responses import StreamingResponse

    name = (file.filename or "purchase.xlsx").strip()
    raw = await file.read()
    if not raw:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Empty file.")
    if not name.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Only .xlsx files are supported for Clear Qty (save .xls as .xlsx first).",
        )
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="openpyxl is not installed on the server.",
        ) from exc

    try:
        wb = load_workbook(BytesIO(raw))
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Invalid Excel file: {exc}") from exc

    if len(wb.worksheets) < 1:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Excel file has no worksheets.")
    ws = wb.worksheets[0]  # Sheet 1
    cleared = 0
    for row in range(1, ws.max_row + 1):
        manual = ws.cell(row=row, column=1).value
        if manual is None or str(manual).strip() == "":
            continue
        # Skip header-ish first row
        if row == 1:
            m = str(manual).strip().lower()
            if not m.replace(".", "", 1).isdigit() and any(k in m for k in ("manual", "item", "id")):
                continue
        # Sheet 1 Column C = Qty
        ws.cell(row=row, column=3).value = 0
        cleared += 1

    out = BytesIO()
    wb.save(out)
    out.seek(0)
    out_name = name.rsplit(".", 1)[0] + "_qty_cleared.xlsx"
    headers = {
        "Content-Disposition": f'attachment; filename="{out_name}"',
        "X-Rows-Cleared": str(cleared),
    }
    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )
